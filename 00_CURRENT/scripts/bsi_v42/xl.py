"""Workbook helpers: v41 house style, formulas with cached values, text lint."""
import re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = 'Sahel'
NAVY = 'FF1F3864'
LIGHT = 'FFD9E2F3'
THIN = Side(style='thin', color='FFBFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F_TITLE = Font(name=FONT, size=14, bold=True, color=NAVY)
F_HEAD = Font(name=FONT, size=10, bold=True, color='FFFFFFFF')
F_BODY = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
FILL_TITLE = PatternFill('solid', fgColor=LIGHT)
FILL_HEAD = PatternFill('solid', fgColor=NAVY)
FILL_WARN = PatternFill('solid', fgColor='FFFFF2CC')
FILL_BAD = PatternFill('solid', fgColor='FFFCE4D6')
FILL_OK = PatternFill('solid', fgColor='FFE2EFDA')
AL = Alignment(horizontal='center', vertical='center', wrap_text=True)
AL_L = Alignment(horizontal='left', vertical='center', wrap_text=True)

DIAC = re.compile('[ً-ٰٟـ]')
DASHES = re.compile('[-‐-―−]')
PERSIAN = re.compile('[؀-ۿ]')

CACHE = {}          # (sheet title, coordinate) -> cached value for formula cells
LINT = []           # (sheet, coordinate, text) violations


class FX:
    """A formula plus the value Python computed for it (stored as the cached value)."""
    def __init__(self, formula, value):
        self.formula = formula
        self.value = value


def lint(sheet, coord, text):
    if not isinstance(text, str) or text.startswith('http') or text.startswith('='):
        return
    if not PERSIAN.search(text):
        return
    if DIAC.search(text) or DASHES.search(text):
        LINT.append((sheet, coord, text[:80]))


class Sheet:
    def __init__(self, wb, title, heading, subtitle, headers, widths=None, freeze=None, raw_cols=()):
        self.ws = wb.create_sheet(title)
        self.title = title
        self.ncols = len(headers)
        self.raw_cols = set(raw_cols)
        ws = self.ws
        ws.sheet_view.rightToLeft = True
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=self.ncols)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=self.ncols)
        c = ws.cell(1, 1, heading); c.font = F_TITLE; c.fill = FILL_TITLE; c.alignment = AL
        c = ws.cell(2, 1, subtitle); c.font = F_BODY; c.alignment = AL
        lint(title, 'A1', heading); lint(title, 'A2', subtitle)
        ws.row_dimensions[1].height = 26
        ws.row_dimensions[2].height = 44
        for j, h in enumerate(headers, 1):
            c = ws.cell(3, j, h); c.font = F_HEAD; c.fill = FILL_HEAD; c.alignment = AL; c.border = BORDER
            lint(title, c.coordinate, h)
        ws.row_dimensions[3].height = 42
        for j in range(1, self.ncols + 1):
            w = (widths or {}).get(j, 16)
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = freeze or 'C4'
        self.r = 4

    def add(self, values, fill=None, height=None):
        ws = self.ws
        assert len(values) == self.ncols, f'{self.title}: {len(values)} values for {self.ncols} columns: {values[:3]}'
        for j, v in enumerate(values, 1):
            c = ws.cell(self.r, j)
            if isinstance(v, FX):
                c.value = v.formula
                CACHE[(self.title, c.coordinate)] = v.value
            else:
                c.value = v
                if j not in self.raw_cols:
                    lint(self.title, c.coordinate, v)
            c.font = F_BODY; c.alignment = AL; c.border = BORDER
            if fill is not None:
                c.fill = fill
        if height:
            ws.row_dimensions[self.r].height = height
        self.r += 1
        return self.r - 1

    def autofilter(self):
        self.ws.auto_filter.ref = f'A3:{get_column_letter(self.ncols)}{self.r - 1}'
