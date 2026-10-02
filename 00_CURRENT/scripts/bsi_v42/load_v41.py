"""Read-only extraction of the data this stage needs from the v41 workbook.

Nothing here modifies v41.  Every value keeps its sheet/cell origin so the
Changes/Measures sheets can cite it.
"""
import openpyxl

MASTER = 'Turnover نهایی'
SEAT = 'Turnover Seat Reconstruction'
AUDIT = 'Turnover Resolution Audit'
SCOPE = 'BR Scope v41'
TENURE = 'BR Tenure v41'
INDEP = 'BR Independence v41'
SOURCES = 'BR Sources v41'
CR1A = 'CR1A v39'

MASTER_COLS = dict(G=7, H=8, J=10, K=11, M=13, N=14, P=16, Q=17, S=19, T=20)


def open_wb(path, data_only=False):
    return openpyxl.load_workbook(path, data_only=data_only)


def read_scope(wb):
    ws = wb[SCOPE]
    out = []
    for r in range(4, ws.max_row + 1):
        if not ws.cell(r, 2).value:
            continue
        out.append(dict(row=r, id=ws.cell(r, 1).value, sym=ws.cell(r, 2).value,
                        legal=ws.cell(r, 3).value, fy_start=ws.cell(r, 4).value,
                        fye=ws.cell(r, 5).value, fye_src=ws.cell(r, 6).value,
                        src_id=ws.cell(r, 7).value, c1=ws.cell(r, 8).value,
                        c2=ws.cell(r, 9).value, fye_status=ws.cell(r, 12).value,
                        note=ws.cell(r, 13).value))
    return out


def read_master(wb):
    ws = wb[MASTER]
    out = {}
    for r in range(4, ws.max_row + 1):
        sym = ws.cell(r, 1).value
        if not sym:
            continue
        d = dict(row=r, sym=sym, group=ws.cell(r, 2).value, quality=ws.cell(r, 3).value,
                 method=ws.cell(r, 4).value, eligible=ws.cell(r, 5).value,
                 legal_struct=ws.cell(r, 6).value, chair_change=ws.cell(r, 22).value,
                 note=ws.cell(r, 23).value, status=ws.cell(r, 24).value,
                 src=ws.cell(r, 26).value)
        for k, c in MASTER_COLS.items():
            d[k] = ws.cell(r, c).value
            d[k + '_cell'] = ws.cell(r, c).coordinate
        out[sym] = d
    return out


def read_seat_recon(wb):
    ws = wb[SEAT]
    rows = []
    for r in range(4, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, 20)]
        if not v[0]:
            continue
        rows.append(dict(row=r, sym=v[0], snapshot=v[2], holder=v[3], rep=v[4], role=v[5],
                         eff=v[6], intro=v[7], cont=v[10], exit=v[11], entry=v[12],
                         change=v[13], src_title=v[14], src_date=v[15], tracing=v[16],
                         url=v[17], note=v[18]))
    return rows


def read_audit(wb):
    ws = wb[AUDIT]
    rows = []
    for r in range(4, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, 23)]
        if not v[0]:
            continue
        rows.append(dict(row=r, sym=v[0], year=v[2], var=v[3], before=v[4], after=v[5],
                         status=v[6], old_member=v[7], new_member=v[8], holder=v[9],
                         old_rep=v[10], new_rep=v[11], change=v[12], eff=v[13],
                         src_title=v[14], pub=v[15], tracing=v[16], url=v[17],
                         ev_base=v[18], ev_end=v[19], conf=v[20], note=v[21]))
    return rows


def read_tenure(wb):
    ws = wb[TENURE]
    rows = []
    for r in range(4, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, 30)]
        if not v[0]:
            continue
        rows.append(dict(row=r, id=v[0], sym=v[1], member_id=v[2], kind=v[3], name=v[4],
                         holder=v[5], idhash=v[6], idbasis=v[7], role=v[8], cls=v[9],
                         start=v[10], lower=v[12], upper=v[14], asof=v[16], datetype=v[22],
                         basis=v[23], report_date=v[24], note=v[28]))
    return rows


def read_roster_v41(wb):
    ws = wb[INDEP]
    rows = []
    for r in range(4, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, 22)]
        if not v[0]:
            continue
        rows.append(dict(row=r, id=v[0], sym=v[1], holder=v[2], name=v[3], kind=v[4],
                         role=v[5], executive=v[6], asof=v[17], ceo=v[19], chair=v[20]))
    return rows


def read_registry(wb):
    ws = wb[SOURCES]
    rows = []
    for r in range(4, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, 12)]
        if not v[0]:
            continue
        rows.append(dict(row=r, sid=v[0], sym=v[1], type=v[2], arch=v[3], title=v[4],
                         tracing=v[5], pub=v[6], url=v[7], role=v[8], sha=v[9], path=v[10]))
    return rows


def read_cr1a(wb):
    ws = wb[CR1A]
    out = {}
    for r in range(4, 34):
        sym = ws.cell(r, 1).value
        if not sym:
            continue
        out[sym] = dict(row=r, status=ws.cell(r, 4).value, conservative=ws.cell(r, 5).value,
                        chair_base=ws.cell(r, 7).value, chair_end=ws.cell(r, 8).value)
    return out
