#!/usr/bin/env python3
"""Independent recalculation of the formulas in the NEW v42 sheets with the `formulas`
library (a pure-Python Excel engine, not LibreOffice) and comparison with the cached values.

usage: python verify_formulas.py <v42.xlsx>
"""
import sys, os, tempfile, json, shutil
import openpyxl, formulas

NEW = ['Scope v42', 'Measures v42', 'Flow Summary v42', 'Coverage v42', 'QA v42']


def main(v42):
    wb = openpyxl.load_workbook(v42)
    keep = NEW
    for n in list(wb.sheetnames):
        if n not in keep:
            del wb[n]
    d = tempfile.mkdtemp()
    p = os.path.join(d, 'new_only.xlsx')
    wb.save(p)
    cached = openpyxl.load_workbook(v42, data_only=True)
    form = openpyxl.load_workbook(v42)
    xl = formulas.ExcelModel().loads(p).finish()
    sol = xl.calculate()
    nform = 0; diffs = []; skipped = 0
    for n in keep:
        for row in form[n].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    nform += 1
                    key = f"'[{os.path.basename(p)}]{n.upper()}'!{c.coordinate}"
                    if key not in sol:
                        skipped += 1
                        continue
                    v = sol[key].value[0, 0]
                    want = cached[n][c.coordinate].value
                    if hasattr(v, 'item'):
                        v = v.item()
                    vv = '' if (v is None or str(v) == 'empty') else v
                    ww = '' if want is None else want
                    if isinstance(vv, (int, float)) and isinstance(ww, (int, float)):
                        ok = abs(vv - ww) < 1e-9
                    else:
                        ok = str(vv) == str(ww)
                    if not ok:
                        diffs.append((n, c.coordinate, str(c.value)[:70], str(vv), str(ww)))
    res = dict(formulas_checked=nform, not_evaluated=skipped, differences=len(diffs), examples=diffs[:10])
    print(json.dumps(res, ensure_ascii=False, indent=1))
    shutil.rmtree(d)
    return res


if __name__ == '__main__':
    main(sys.argv[1])
