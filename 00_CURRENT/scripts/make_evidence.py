#!/usr/bin/env python3
"""Run the whole stage-3 pipeline and write the deliverables.

usage: python make_evidence.py <v41.xlsx> <00_CURRENT dir>
Outputs (in <00_CURRENT dir>):
  Iran_BSI_v42_Turnover_Person_Seat_Reconciliation.xlsx
  Iran_BSI_v42_Turnover_Person_Seat_Evidence.zip
  Iran_BSI_v42_Data_Dictionary.md
"""
import sys, os, io, csv, json, shutil, subprocess, zipfile, datetime, hashlib, re
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import openpyxl
import build_v42 as B
import verify_output, verify_formulas, verify_independent, review_pass
from bsi_v42 import load_v41 as L

ZIP_NAME = 'Iran_BSI_v42_Turnover_Person_Seat_Evidence.zip'
DICT_NAME = 'Iran_BSI_v42_Data_Dictionary.md'
V41_NAME = 'Iran_BSI_v41_Independence_Roles_Tenure.xlsx'


def sheet_csv(ws):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator='\n')
    for row in ws.iter_rows(values_only=True):
        w.writerow(['' if v is None else v for v in row])
    return buf.getvalue()


def fetch_log(registry):
    lines = [f'# fetch attempts, run {datetime.datetime.utcnow().isoformat()}Z (outbound proxy of the cloud session)']
    urls = [r['url'] for r in registry if r['tracing'] in (1444154, 1443078, 1317785)][:3]
    urls += ['https://www.codal.ir/', 'https://search.codal.ir/api/search/v2/q?PageNumber=1', 'https://www.irisl.net/']
    for u in urls:
        try:
            p = subprocess.run(['curl', '-sS', '-m', '15', '-o', '/dev/null', '-w', '%{http_code}', u], capture_output=True, text=True, timeout=30)
            lines.append(f'{u[:110]} -> http={p.stdout.strip() or "none"} stderr={p.stderr.strip()[:90]}')
        except Exception as e:                                     # noqa
            lines.append(f'{u[:110]} -> exception {type(e).__name__}')
    lines.append('WebFetch tool on a Codal Decision.aspx page (tracing 1444154) in the same session -> HTTP 503 Service Unavailable')
    lines.append('Meaning: documents could not be re-opened in this stage. This is not evidence that any document does not exist.')
    return '\n'.join(lines) + '\n'


def dictionary_md(wb):
    ws = wb['Dictionary v42']
    out = ['# فرهنگ داده v42', '', 'خروجی برگه Dictionary v42 از فایل اکسل.', '']
    for r in ws.iter_rows(min_row=4, values_only=True):
        if not r[0]:
            continue
        out += [f'## {r[0]}: {r[1]}', f'- وضعیت: {r[2]}', f'- تعریف: {r[3]}', f'- واحد: {r[4]}', f'- منبع داده: {r[5]}',
                f'- محدودیت و تفاوت با معیار دیگر: {r[6]}', f'- خانواده: {r[7]}؛ نوع: {r[8]}', '']
    return '\n'.join(out)


def request_list(m, wb):
    rows = [['شرکت', 'مورد باز', 'سند مشخص لازم', 'شماره ردیابی یا منبع کاندید', 'برگه مرتبط']]
    for c in wb['Five Cases v42'].iter_rows(min_row=4, values_only=True):
        if c[0] and c[0] != 'SUM' and c[4] in ('داده کافی پیدا نشد', 'منبع در دسترس نبود', 'متعارض', 'هنوز بررسی نشده', 'تایید جزئی') and c[7]:
            rows.append([c[1], f'{c[0]}: {c[2]}', c[7], c[6], 'Five Cases v42'])
    for s in m['scope']:
        sym = s['sym']
        if m['tier'][sym] in ('B_AGG', 'B_CONTRA'):
            tb = m['tb'][sym]
            bn, en = tb['base_notice'], tb['end_notice']
            rows.append([sym, 'فهرست اسمی مبنا و پایان برای ورود و خروج افراد یکتا',
                         'متن اطلاعیه ن۴۵ نزدیک مبنا و پایان و جدول گزارش FY1403 و FY1404',
                         f"مبنا: {bn['tracing'] if bn else 'ندارد'}؛ پایان: {en['tracing'] if en else 'ندارد'}؛ گزارش FY1404: {s['src_id']}", 'Scope v42'])
    return rows


def main(v41, out):
    os.makedirs(out, exist_ok=True)
    shutil.copy2(v41, os.path.join(out, V41_NAME))
    m, ctx, info = B.main(v41, out)
    final = info['final']
    nd = B.count_old_diffs(v41, final)
    if nd:
        m, ctx, info = B.main(v41, out, master_changed=nd)
    logs = {}
    r1 = verify_output.main(v41, final); logs['verify_output.json'] = r1
    r2 = verify_formulas.main(final); logs['verify_formulas.json'] = r2
    r3 = verify_independent.main(final); logs['verify_independent.json'] = r3
    r4 = review_pass.main(v41, final); logs['review_pass.json'] = r4
    tests = subprocess.run([sys.executable, '-m', 'pytest', '-q', os.path.join(HERE, 'tests')], capture_output=True, text=True, cwd=HERE)
    wb = openpyxl.load_workbook(final, data_only=True)
    with open(os.path.join(out, DICT_NAME), 'w', encoding='utf-8') as f:
        f.write(dictionary_md(wb))
    zpath = os.path.join(out, ZIP_NAME)
    with zipfile.ZipFile(zpath, 'w', zipfile.ZIP_DEFLATED) as z:
        spec = info['spec41']; spec['file'] = V41_NAME
        z.writestr('inputs/v41_spec.json', json.dumps(spec, ensure_ascii=False, indent=1))
        reg = L.read_registry(L.open_wb(v41))
        buf = io.StringIO(); w = csv.writer(buf, lineterminator='\n')
        w.writerow(['Source_ID', 'symbol', 'type', 'archive_id', 'title', 'tracing', 'published', 'url', 'role', 'sha256', 'local_path_in_v41_zip'])
        for r in reg:
            w.writerow([r['sid'], r['sym'], r['type'], r['arch'], r['title'], r['tracing'], r['pub'], r['url'], r['role'], r['sha'], r['path']])
        z.writestr('inputs/registry_v41_extract.csv', buf.getvalue())
        for n in info['new_sheets']:
            z.writestr('tables/' + n.replace(' ', '_') + '.csv', sheet_csv(wb[n]))
        rl = request_list(m, wb)
        buf = io.StringIO(); csv.writer(buf, lineterminator='\n').writerows(rl)
        z.writestr('requests/source_request_list.csv', buf.getvalue())
        for name, data in logs.items():
            z.writestr('logs/' + name, json.dumps(data, ensure_ascii=False, indent=1, default=str))
        z.writestr('logs/test_results.txt', tests.stdout + tests.stderr)
        z.writestr('logs/fetch_attempts.txt', fetch_log(reg))
        z.writestr('logs/build_summary.json', json.dumps({k: v for k, v in info.items() if k != 'spec41'}, ensure_ascii=False, indent=1))
        for root, dirs, files in os.walk(HERE):
            dirs[:] = [d for d in dirs if d not in ('__pycache__', '.pytest_cache')]
            for fn in files:
                if fn.endswith('.pyc'):
                    continue
                p = os.path.join(root, fn)
                z.write(p, 'scripts/' + os.path.relpath(p, HERE))
        z.write(os.path.join(out, DICT_NAME), 'dictionary/' + DICT_NAME)
        z.writestr('README_v42.md', README.format(sha=spec['sha256'], sheets=spec['n_sheets']))
    print('pytest:', tests.stdout.strip().splitlines()[-1] if tests.stdout.strip() else tests.stderr[:200])
    print('wrote', final, zpath)


README = """# بسته شواهد v42 (مرحله سوم از پنج مرحله)

فایل مبنا: Iran_BSI_v41_Independence_Roles_Tenure.xlsx
sha256 فایل مبنا: {sha}
تعداد برگه های v41: {sheets} (در v42 بدون تغییر)

## محتوا
- inputs: مشخصات و هش v41 و استخراج رجیستری منابع v41 (438 ردیف با sha256 و مسیر محلی در بسته v41)
- tables: خروجی CSV هر برگه جدید v42
- requests: فهرست سند لازم برای هر مورد باز
- logs: آزمون ها، بازمحاسبه مستقل، کنترل فرمول ها، و تلاش های دریافت از کدال
- scripts: کد بازتولید (build_v42.py، verify_*.py، make_evidence.py، پوشه bsi_v42 و tests)
- dictionary: فرهنگ داده

## محدودیت ها (صریح)
- بسته شواهد v41 (Iran_BSI_v41_Independence_Roles_Tenure_Evidence.zip) به این جلسه داده نشد؛ فایل های HTML و PDF اسناد کدال در اختیار نبود.
- کدال و WebFetch از محیط اجرا پاسخ 503 یا قطع اتصال دادند؛ هیچ سندی بازدریافت نشد.
- بنابراین همه وضعیت ها از متن و بازسازی داخل خود v41 (شامل بازسازی v38) آمده و تایید مستقل متن اسناد نیست.
- بازبینی انجام شده بازبینی هوش مصنوعی است و ارزیابی مستقل انسانی نیست.

## بازتولید
pip install -r scripts/requirements.txt
python scripts/make_evidence.py <مسیر v41.xlsx> <پوشه خروجی>
"""

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
